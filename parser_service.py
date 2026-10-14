from pathlib import Path
from docx import Document
from pypdf import PdfReader
from logger import log

class ParserService:
    def __init__(self, folder_path: str | Path, log_file_path: str | Path):
        self.target_folder = Path(folder_path)
        self.history_file = Path(log_file_path)
        self.supported_extensions = {".txt", ".pdf", ".docx"}

    def read_text_file(self, path: Path) -> str:
        with open(path, "r", encoding="utf-8") as f:
            return f.read()

    def read_pdf_file(self, path: Path) -> str:
        reader = PdfReader(path)
        text_content = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_content.append(text)
        return "\n".join(text_content)

    def read_docx_file(self, path: Path) -> str:
        doc = Document(path)
        text_content = []

        for p in doc.paragraphs:
            if p.text.strip():
                text_content.append(p.text)

        for table_index, table in enumerate(doc.tables, 1):
            text_content.append(f"\n--- [Table {table_index} Content] ---")
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text and (not row_text or row_text[-1] != cell_text):
                        row_text.append(cell_text)
                if row_text:
                    text_content.append(" | ".join(row_text))

        return "\n".join(text_content)

    def process_file_router(self, path: Path) -> str:
        extension = path.suffix.lower()
        if extension == ".txt":
            return self.read_text_file(path)
        elif extension == ".pdf":
            return self.read_pdf_file(path)
        elif extension == ".docx":
            return self.read_docx_file(path)
        else:
            raise ValueError(f"Unsupported file type: {extension}")

    def process_folder(self) -> list[str]:
        processed_files = set()
        all_new_contents = [] # <-- 1. Create a container to hold texts

        if self.history_file.exists():
            with open(self.history_file, "r", encoding="utf-8") as hf:
                processed_files = {line.strip() for line in hf if line.strip()}

        log(f"Loaded {len(processed_files)} previously processed files.")

        if not self.target_folder.exists():
            # Help catch path setup mistakes immediately
            raise FileNotFoundError(f"Directory missing at: {self.target_folder.resolve()}")

        for file_path in self.target_folder.iterdir():
            if file_path.is_dir() or file_path.suffix.lower() not in self.supported_extensions:
                continue

            file_identifier = str(file_path.resolve())

            if file_identifier in processed_files:
                log(f"Skipping: {file_path.name} (Already processed)")
                continue

            log(f"\nProcessing: {file_path.name}...")

            try:
                file_content = self.process_file_router(file_path)
                
                # 2. Collect the extracted text
                all_new_contents.append(file_content)

                with open(self.history_file, "a", encoding="utf-8") as hf:
                    hf.write(file_identifier + "\n")

                processed_files.add(file_identifier)

            except Exception as e:
                log(f"❌ Failed to process {file_path.name}. Error: {e}")
                
        # 3. Return the list of strings back to your main ResumeParser script
        return all_new_contents

